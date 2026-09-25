import csv
import io
import random
from datetime import datetime, timezone
from typing import Dict, Any
from fastapi import APIRouter, Depends, HTTPException, Query, status, WebSocket, WebSocketDisconnect
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session
from auth_dep import require_admin, is_valid_admin_token
from database import get_db
from models import Order, OrderItem, TokenRecord
from realtime import order_manager
from schemas import OrderCreate, OrderStatusUpdate

router = APIRouter(prefix="/orders", tags=["Orders Management"])


@router.websocket("/ws")
async def websocket_orders_endpoint(
    websocket: WebSocket,
    token: str = "",
    db: Session = Depends(get_db),
):
    # Kasir realtime: tolak koneksi tanpa token admin yang valid.
    if not is_valid_admin_token(db, token):
        await websocket.close(code=4401)
        return
    await order_manager.connect(websocket)
    try:
        while True:
            data = await websocket.receive_text()
            if data == "ping":
                await websocket.send_text("pong")
    except WebSocketDisconnect:
        order_manager.disconnect(websocket)
    except Exception:
        order_manager.disconnect(websocket)


def serialize_order(ord: Order) -> Dict[str, Any]:
    return {
        "orderId": ord.order_id,
        "createdAt": ord.created_at.isoformat() if ord.created_at else "",
        "customer": {
            "name": ord.customer_name,
            "phone": ord.customer_phone,
            "token": ord.customer_token,
            "orderType": ord.order_type,
            "tableOrAddress": ord.table_or_address,
            "specialRequest": ord.special_request,
        },
        "payment": {
            "method": ord.payment_method,
            "label": ord.payment_label,
            "reference": ord.payment_reference,
            "status": ord.payment_status,
            "paid": bool(ord.is_paid),
        },
        "items": [
            {
                "id": it.menu_item_id or str(it.id),
                "name": it.name,
                "price": it.price,
                "quantity": it.quantity,
                "notes": it.notes or "",
                "subtotal": it.subtotal,
            }
            for it in ord.items
        ],
        "breakdown": {
            "subtotal": ord.subtotal,
            "ecoFee": ord.eco_fee,
            "tax": ord.tax,
            "total": ord.total,
        },
        "orderStatus": ord.order_status,
    }


@router.get("")
def get_all_orders(db: Session = Depends(get_db), _admin: bool = Depends(require_admin)):
    """Retrieve all orders sorted by createdAt descending (Admin)."""
    orders = db.query(Order).order_by(Order.created_at.desc()).all()
    return [serialize_order(o) for o in orders]


EXPORT_COLUMNS = [
    "No Pesanan", "Waktu", "Pelanggan", "Telepon", "Token",
    "Tipe Layanan", "Meja/Alamat", "Menu", "Qty", "Harga Satuan",
    "Subtotal Item", "Total Order", "Metode Bayar", "Status Bayar", "Status Order",
]


def _export_rows(db: Session):
    """One row per order item (orders without items yield one row with empty menu)."""
    orders = db.query(Order).order_by(Order.created_at.desc()).all()
    rows = []
    for ord in orders:
        base = [
            ord.order_id,
            ord.created_at.isoformat() if ord.created_at else "",
            ord.customer_name,
            ord.customer_phone,
            ord.customer_token or "-",
            ord.order_type,
            ord.table_or_address,
            "", 0, 0.0, 0.0,
            ord.total,
            ord.payment_label or ord.payment_method,
            ord.payment_status,
            ord.order_status,
        ]
        if ord.items:
            for it in ord.items:
                r = list(base)
                r[7] = it.name
                r[8] = it.quantity
                r[9] = it.price
                r[10] = it.subtotal
                rows.append(r)
        else:
            rows.append(base)
    return rows


@router.get("/stats")
def get_sales_stats(db: Session = Depends(get_db), _admin: bool = Depends(require_admin)):
    """Aggregated sales analytics (Admin)."""
    total_orders = db.query(func.count(Order.id)).scalar() or 0
    total_revenue = db.query(func.coalesce(func.sum(Order.total), 0.0)).scalar() or 0.0
    total_portions = db.query(func.coalesce(func.sum(OrderItem.quantity), 0)).scalar() or 0

    top_menus_q = (
        db.query(
            OrderItem.name.label("name"),
            func.sum(OrderItem.quantity).label("qty"),
            func.sum(OrderItem.subtotal).label("revenue"),
            func.count(func.distinct(OrderItem.order_id)).label("orders"),
        )
        .group_by(OrderItem.name)
        .order_by(func.sum(OrderItem.quantity).desc())
        .limit(10)
        .all()
    )
    top_menus = [
        {"name": r.name, "qty": int(r.qty or 0), "revenue": float(r.revenue or 0.0), "orders": int(r.orders or 0)}
        for r in top_menus_q
    ]

    pay_q = (
        db.query(
            Order.payment_method.label("method"),
            func.count(Order.id).label("count"),
            func.coalesce(func.sum(Order.total), 0.0).label("revenue"),
        )
        .group_by(Order.payment_method)
        .all()
    )
    payment_methods = []
    for r in pay_q:
        method = (r.method or "unknown").lower()
        label = "QRIS" if method == "qris" else ("Tunai di Kasir" if method in ("cash", "tunai") else (r.method or "-"))
        count = int(r.count or 0)
        payment_methods.append({
            "method": method,
            "label": label,
            "count": count,
            "pct": round(count / total_orders * 100, 1) if total_orders else 0.0,
            "revenue": float(r.revenue or 0.0),
        })
    payment_methods.sort(key=lambda x: x["count"], reverse=True)

    return {
        "summary": {
            "total_orders": total_orders,
            "total_revenue": float(total_revenue),
            "total_portions": int(total_portions),
        },
        "top_menus": top_menus,
        "payment_methods": payment_methods,
    }


@router.get("/export")
def export_sales_data(format: str = Query(default="xlsx", pattern="^(xlsx|csv)$"), db: Session = Depends(get_db), _admin: bool = Depends(require_admin)):
    """Export sales data (Admin). ?format=xlsx (default) or ?format=csv. One row per item."""
    rows = _export_rows(db)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")

    if format == "csv":
        buf = io.StringIO()
        buf.write("\ufeff")  # BOM so Excel opens Indonesian text correctly
        writer = csv.writer(buf)
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows(rows)
        data = buf.getvalue().encode("utf-8")
        return StreamingResponse(
            io.BytesIO(data),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": f"attachment; filename=penjualan-sikopi-{stamp}.csv"},
        )

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Penjualan"
    ws.append(EXPORT_COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    widths = [12, 20, 18, 12, 8, 14, 14, 28, 6, 12, 13, 12, 14, 18, 18]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=penjualan-sikopi-{stamp}.xlsx"},
    )


@router.get("/{order_id}")
def get_order_by_id(order_id: str, db: Session = Depends(get_db)):
    """Retrieve a single order by orderId."""
    ord = db.query(Order).filter(Order.order_id == order_id).first()
    if not ord:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan.")
    return serialize_order(ord)


@router.post("", status_code=status.HTTP_201_CREATED)
async def create_order(payload: OrderCreate, db: Session = Depends(get_db)):
    """Place a new order, calculate totals, and burn the customer token."""
    # 1. Generate Order Number
    seq = random.randint(10000, 99900)
    order_id = f"HYT-2026-{seq}"
    now = datetime.now(timezone.utc)

    # 2. Calculate Subtotal from Items
    items_data = []
    subtotal = 0.0

    for it in payload.items:
        it_subtotal = it.price * it.quantity
        subtotal += it_subtotal
        items_data.append({
            "menu_item_id": it.id,
            "name": it.name,
            "price": it.price,
            "quantity": it.quantity,
            "notes": it.notes or "",
            "subtotal": it_subtotal
        })

    # Mode Pengujian: Biaya tambahan 0 sehingga total murni akumulasi item
    eco_fee = 0.0
    tax = 0.0
    total = subtotal

    if payload.breakdown:
        subtotal = payload.breakdown.subtotal
        eco_fee = payload.breakdown.ecoFee
        tax = payload.breakdown.tax
        total = payload.breakdown.total

    # 3. Create Order
    is_qris = payload.payment.method.lower() == "qris"
    new_order = Order(
        order_id=order_id,
        created_at=now,
        customer_name=payload.customer.name.strip() or "Pelanggan",
        customer_phone=payload.customer.phone or "-",
        customer_token=payload.customer.token or None,
        order_type=payload.customer.orderType or "Makan di Tempat",
        table_or_address=payload.customer.tableOrAddress or "Meja Reguler",
        special_request=payload.customer.specialRequest or "-",
        payment_method=payload.payment.method,
        payment_label=payload.payment.label or ("QRIS" if is_qris else "Bayar di Kasir (Tunai / EDC)"),
        payment_reference=payload.payment.reference or f"REF-{random.randint(100000, 999000)}",
        payment_status="Menunggu Pembayaran",
        subtotal=subtotal,
        eco_fee=eco_fee,
        tax=tax,
        total=total,
        order_status="Diterima & Disiapkan"
    )
    db.add(new_order)
    db.flush()

    # 4. Insert Order Items
    for item_dict in items_data:
        order_item = OrderItem(
            order_id=order_id,
            menu_item_id=item_dict["menu_item_id"],
            name=item_dict["name"],
            price=item_dict["price"],
            quantity=item_dict["quantity"],
            notes=item_dict["notes"],
            subtotal=item_dict["subtotal"]
        )
        db.add(order_item)

    # 5. Burn Customer Token (One-time use strictly enforced: reject reuse)
    cust_token = payload.customer.token
    if cust_token:
        clean_token = str(cust_token).strip()
        existing_token = db.query(TokenRecord).filter(TokenRecord.token == clean_token).first()
        if existing_token and existing_token.is_used:
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail=f'Token "{clean_token}" sudah dipakai. Minta token baru ke kasir.'
            )
        if existing_token:
            existing_token.is_used = True
            existing_token.order_id = order_id
            existing_token.customer_name = payload.customer.name
            existing_token.burned_at = now
        else:
            db.add(TokenRecord(
                token=clean_token,
                is_used=True,
                customer_name=payload.customer.name,
                order_id=order_id,
                burned_at=now
            ))

    db.commit()
    db.refresh(new_order)
    serialized = serialize_order(new_order)
    await order_manager.broadcast("NEW_ORDER", serialized)
    return serialized


@router.patch("/{order_id}/status")
async def update_order_status(order_id: str, payload: OrderStatusUpdate, db: Session = Depends(get_db), _admin: bool = Depends(require_admin)):
    """Update order lifecycle status or payment status (Admin)."""
    ord = db.query(Order).filter(Order.order_id == order_id).first()
    if not ord:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan.")

    if payload.order_status:
        ord.order_status = payload.order_status
    if payload.is_paid is not None:
        ord.is_paid = payload.is_paid
        ord.payment_status = "Lunas" if payload.is_paid else "Menunggu Pembayaran"
    elif payload.payment_status:
        ord.payment_status = payload.payment_status
        if "lunas" in payload.payment_status.lower():
            ord.is_paid = True

    db.commit()
    db.refresh(ord)
    serialized = serialize_order(ord)
    await order_manager.broadcast("STATUS_UPDATED", serialized)
    return serialized


@router.delete("/{order_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_order(order_id: str, db: Session = Depends(get_db), _admin: bool = Depends(require_admin)):
    """Delete an order (Admin)."""
    ord = db.query(Order).filter(Order.order_id == order_id).first()
    if not ord:
        raise HTTPException(status_code=404, detail="Pesanan tidak ditemukan.")

    db.delete(ord)
    db.commit()
    return None
